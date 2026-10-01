import io
from datetime import date, datetime
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from typing import Optional, Dict, Any
from ..db import db
from ..security import require, has_perm

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

router = APIRouter(prefix="/api/reports", tags=["reports-export"])

NAVY = colors.HexColor("#0B1F3A")
GOLD = colors.HexColor("#C9A227")
LIGHT = colors.HexColor("#F1F5F9")


def _settings_title():
    return "INFORMASI KEPEGAWAIAN LAGUSIT"


def _pdf(title: str, subtitle: str, headers, rows, landscape_mode=False):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4) if landscape_mode else A4,
                            topMargin=18 * mm, bottomMargin=16 * mm, leftMargin=16 * mm, rightMargin=16 * mm)
    styles = getSampleStyleSheet()
    h = ParagraphStyle("h", parent=styles["Title"], fontSize=15, textColor=NAVY, spaceAfter=2)
    sub = ParagraphStyle("sub", parent=styles["Normal"], fontSize=10, textColor=colors.HexColor("#64748B"), spaceAfter=10)
    org = ParagraphStyle("org", parent=styles["Normal"], fontSize=9, textColor=GOLD, spaceAfter=0)
    cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=8, leading=10)
    elems = [Paragraph("LAPAS GUNUNGSITOLI", org), Paragraph(title, h), Paragraph(subtitle, sub)]
    data = [[Paragraph(f"<b>{c}</b>", ParagraphStyle("hc", parent=cell, textColor=colors.white)) for c in headers]]
    for r in rows:
        data.append([Paragraph(str(v) if v is not None else "-", cell) for v in r])
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))
    elems.append(t)
    elems.append(Spacer(1, 10))
    elems.append(Paragraph(f"Dicetak: {datetime.now().strftime('%d-%m-%Y %H:%M')}",
                           ParagraphStyle("f", parent=styles["Normal"], fontSize=7, textColor=colors.grey)))
    doc.build(elems)
    buf.seek(0)
    return buf


def _excel(title: str, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan"
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(headers)))
    ws["A1"].font = Font(bold=True, size=13, color="0B1F3A")
    ws.append(list(headers))
    hdr_fill = PatternFill("solid", fgColor="0B1F3A")
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(list(r))
    from openpyxl.utils import get_column_letter
    for i in range(1, len(headers) + 1):
        letter = get_column_letter(i)
        width = 10
        for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            v = row[0].value
            if v is not None:
                width = max(width, len(str(v)))
        ws.column_dimensions[letter].width = min(width + 4, 45)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _stream(buf, media, filename):
    return StreamingResponse(buf, media_type=media, headers={"Content-Disposition": f'attachment; filename="{filename}"'})


# ---------------- Employee report ----------------
@router.get("/employees/export")
async def export_employees(format: str = "excel", user: dict = Depends(require("report.export", "employee.view_all"))):
    headers = ["NIP", "Nama", "Jabatan", "Unit", "Golongan", "Status"]
    rows = []
    async for e in db.employees.find({"deleted": {"$ne": True}}, {"_id": 0}).sort("name", 1):
        async def nm(coll, _id):
            if not _id:
                return "-"
            d = await db[coll].find_one({"id": _id}, {"_id": 0, "name": 1})
            return d["name"] if d else "-"
        g = await db.grades.find_one({"id": e.get("grade_id")}, {"_id": 0}) if e.get("grade_id") else None
        rows.append([e.get("nip"), e.get("name"), await nm("positions", e.get("position_id")),
                     await nm("organizational_units", e.get("unit_id")), g["name"] if g else "-", e.get("status_kepegawaian")])
    title = "Laporan Data Pegawai"
    if format == "pdf":
        return _stream(_pdf(title, f"Total {len(rows)} pegawai", headers, rows), "application/pdf", "data-pegawai.pdf")
    return _stream(_excel(title, headers, rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "data-pegawai.xlsx")


# ---------------- Shared helpers ----------------
async def _emp(eid):
    if not eid:
        return {"name": "-", "nip": "-"}
    e = await db.employees.find_one({"id": eid}, {"_id": 0, "name": 1, "nip": 1})
    return e or {"name": "-", "nip": "-"}


def _due_status(next_date, thr):
    try:
        dl = (datetime.strptime((next_date or "")[:10], "%Y-%m-%d").date() - date.today()).days
    except (ValueError, TypeError):
        return "-", None
    if dl < 0:
        return "Terlambat", dl
    if dl <= 30:
        return "Jatuh Tempo", dl
    if dl <= thr:
        return "Mendekati", dl
    return "Aman", dl


async def _out(format, title, sub, headers, rows, fname, landscape_mode=False):
    if format == "pdf":
        return _stream(_pdf(title, sub, headers, rows, landscape_mode=landscape_mode), "application/pdf", f"{fname}.pdf")
    return _stream(_excel(title, headers, rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"{fname}.xlsx")


# ---------------- KGB report ----------------
@router.get("/kgb/export")
async def export_kgb(format: str = "excel", user: dict = Depends(require("report.export", "kgb.view"))):
    s = await db.system_settings.find_one({"id": "global"}) or {}
    thr = s.get("kgb_threshold_days", 90)
    headers = ["Nama", "NIP", "Nomor SK", "Tanggal Terakhir", "KGB Berikutnya", "Sisa Hari", "Status"]
    rows = []
    async for k in db.kgb_records.find({"deleted": {"$ne": True}}).sort("next_date", 1):
        e = await _emp(k.get("employee_id"))
        st, dl = _due_status(k.get("next_date"), thr)
        rows.append([e["name"], e["nip"], k.get("nomor_sk"), (k.get("last_date") or "-")[:10], (k.get("next_date") or "-")[:10], f"{dl} hari" if dl is not None else "-", st])
    return await _out(format, "Laporan KGB", f"Total {len(rows)} data", headers, rows, "laporan-kgb")


# ---------------- Promotion report ----------------
@router.get("/promotion/export")
async def export_promotion(format: str = "excel", user: dict = Depends(require("report.export", "promotion.view"))):
    s = await db.system_settings.find_one({"id": "global"}) or {}
    thr = s.get("promotion_threshold_days", 120)
    headers = ["Nama", "NIP", "Pangkat Sekarang", "Pangkat Tujuan", "Periode Berikutnya", "Sisa Hari", "Status"]
    rows = []
    async for p in db.promotion_records.find({"deleted": {"$ne": True}}).sort("next_date", 1):
        e = await _emp(p.get("employee_id"))
        st, dl = _due_status(p.get("next_date"), thr)
        rows.append([e["name"], e["nip"], p.get("pangkat_sekarang"), p.get("pangkat_tujuan"), (p.get("next_date") or "-")[:10], f"{dl} hari" if dl is not None else "-", st])
    return await _out(format, "Laporan Kenaikan Pangkat", f"Total {len(rows)} data", headers, rows, "laporan-kenaikan-pangkat")


# ---------------- Leave report ----------------
@router.get("/leave/export")
async def export_leave(format: str = "excel", user: dict = Depends(require("report.export", "leave.view"))):
    headers = ["Nama", "NIP", "Jenis Cuti", "Mulai", "Selesai", "Status", "Alasan"]
    rows = []
    async for l in db.leave_records.find({"deleted": {"$ne": True}}).sort("created_at", -1):
        e = await _emp(l.get("employee_id"))
        rows.append([e["name"], e["nip"], l.get("jenis"), (l.get("tanggal_mulai") or "-")[:10], (l.get("tanggal_selesai") or "-")[:10], l.get("status"), l.get("alasan")])
    return await _out(format, "Laporan Cuti", f"Total {len(rows)} data", headers, rows, "laporan-cuti")


# ---------------- Evaluation report ----------------
@router.get("/evaluation/export")
async def export_evaluation(format: str = "excel", user: dict = Depends(require("report.export", "evaluation.view"))):
    headers = ["Nama", "Periode", "Nilai Manual", "Nilai Kehadiran (%)", "Penilai", "Status"]
    rows = []
    async for ev in db.evaluations.find({"deleted": {"$ne": True}}).sort("created_at", -1):
        e = await _emp(ev.get("employee_id"))
        rows.append([e["name"], ev.get("period"), ev.get("manual_score"), ev.get("attendance_percentage"), ev.get("evaluator_name"), ev.get("status")])
    return await _out(format, "Laporan Penilaian Pegawai", f"Total {len(rows)} data", headers, rows, "laporan-penilaian")


# ---------------- Notula Apel report (filterable) ----------------
async def _notula_data(date_from=None, date_to=None, pembina=None, location=None, attendance_type_id=None):
    from .attendance import compute_event_recap
    q: Dict[str, Any] = {"deleted": {"$ne": True}}
    if attendance_type_id:
        q["attendance_type_id"] = attendance_type_id
    else:
        q["attendance_type_id"] = "att-staf"  # Notula is Apel Staf by concept
    if date_from or date_to:
        q["date"] = {}
        if date_from:
            q["date"]["$gte"] = date_from
        if date_to:
            q["date"]["$lte"] = date_to
    if pembina:
        q["pembina"] = {"$regex": pembina, "$options": "i"}
    if location:
        q["location"] = {"$regex": location, "$options": "i"}
    types = {t["id"]: t["name"] async for t in db.attendance_types.find({})}
    out = []
    async for ev in db.attendance_events.find(q, {"_id": 0}).sort("date", -1):
        note = await db.attendance_notes.find_one({"event_id": ev["id"]}, {"_id": 0}) or {}
        recap = await compute_event_recap(ev["id"])
        out.append({
            "code": ev.get("code"), "date": ev.get("date"), "time": ev.get("time"),
            "type_name": types.get(ev.get("attendance_type_id"), "-"), "session": ev.get("session"),
            "location": ev.get("location"), "pembina": ev.get("pembina"), "pembina_jabatan": ev.get("pembina_jabatan"),
            "recap": recap,
            "tema": note.get("tema"), "pokok": note.get("pokok"), "isi": note.get("isi"),
            "tindak_lanjut": note.get("tindak_lanjut"), "catatan": note.get("catatan"),
        })
    return out


@router.get("/notula")
async def notula_report(date_from: Optional[str] = None, date_to: Optional[str] = None,
                        pembina: Optional[str] = None, location: Optional[str] = None,
                        attendance_type_id: Optional[str] = None,
                        user: dict = Depends(require("report.view", "attendance.view"))):
    return {"items": await _notula_data(date_from, date_to, pembina, location, attendance_type_id)}


@router.get("/notula/export")
async def notula_export(date_from: Optional[str] = None, date_to: Optional[str] = None,
                        pembina: Optional[str] = None, location: Optional[str] = None,
                        attendance_type_id: Optional[str] = None,
                        user: dict = Depends(require("report.export", "attendance.export"))):
    items = await _notula_data(date_from, date_to, pembina, location, attendance_type_id)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=16 * mm, leftMargin=16 * mm, rightMargin=16 * mm)
    styles = getSampleStyleSheet()
    org = ParagraphStyle("org", parent=styles["Normal"], fontSize=9, textColor=GOLD)
    h = ParagraphStyle("h", parent=styles["Title"], fontSize=14, textColor=NAVY, spaceAfter=2)
    sub = ParagraphStyle("sub", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#64748B"), spaceAfter=8)
    sec = ParagraphStyle("sec", parent=styles["Heading2"], fontSize=11, textColor=NAVY, spaceBefore=12, spaceAfter=3)
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=9, leading=13)
    period = f"Periode: {date_from or 'awal'} s/d {date_to or 'sekarang'}"
    elems = [Paragraph("LAPAS GUNUNGSITOLI", org), Paragraph("LAPORAN NOTULA APEL STAF", h), Paragraph(period, sub)]
    if not items:
        elems.append(Paragraph("Tidak ada data notula pada filter ini.", body))
    for it in items:
        r = it["recap"]
        elems.append(Paragraph(f"{it['code']} — {it['date']} {it['time']} WIB · {it['session']}", sec))
        elems.append(Paragraph(f"<b>Lokasi:</b> {it['location'] or '-'} &nbsp;&nbsp; <b>Pembina:</b> {it['pembina'] or '-'} ({it['pembina_jabatan'] or '-'})", body))
        elems.append(Paragraph(f"<b>Rekap:</b> Hadir {r['present']}/{r['total']} ({r['percentage']}%)", body))
        elems.append(Paragraph(f"<b>Tema:</b> {it['tema'] or '-'}", body))
        elems.append(Paragraph(f"<b>Pokok Amanat:</b> {it['pokok'] or '-'}", body))
        elems.append(Paragraph(f"<b>Ringkasan:</b> {it['isi'] or '-'}", body))
        elems.append(Paragraph(f"<b>Tindak Lanjut:</b> {it['tindak_lanjut'] or '-'}", body))
        elems.append(Paragraph(f"<b>Catatan:</b> {it['catatan'] or '-'}", body))
    elems.append(Spacer(1, 10))
    elems.append(Paragraph(f"Dicetak: {datetime.now().strftime('%d-%m-%Y %H:%M')}", ParagraphStyle("f", parent=body, fontSize=7, textColor=colors.grey)))
    doc.build(elems)
    buf.seek(0)
    return _stream(buf, "application/pdf", "laporan-notula-apel.pdf")


# ---------------- Attendance recap report ----------------
async def _recap_rows(attendance_type_id=None, date_from=None, date_to=None):
    statuses = {s["id"]: s async for s in db.attendance_statuses.find({"deleted": {"$ne": True}})}
    ev_query: Dict[str, Any] = {"deleted": {"$ne": True}}
    if attendance_type_id:
        ev_query["attendance_type_id"] = attendance_type_id
    if date_from or date_to:
        ev_query["date"] = {}
        if date_from:
            ev_query["date"]["$gte"] = date_from
        if date_to:
            ev_query["date"]["$lte"] = date_to
    event_ids = [e["id"] async for e in db.attendance_events.find(ev_query, {"id": 1})]
    agg: Dict[str, Dict[str, Any]] = {}
    async for rec in db.attendance_records.find({"event_id": {"$in": event_ids}}):
        tid = rec.get("attendance_type_id")
        agg.setdefault(tid, {"events": set(), "by_key": {}})
        agg[tid]["events"].add(rec.get("event_id"))
        st = statuses.get(rec.get("status_id"))
        key = st["key"] if st else rec.get("status_id")
        agg[tid]["by_key"][key] = agg[tid]["by_key"].get(key, 0) + 1
    types = {t["id"]: t["name"] async for t in db.attendance_types.find({})}
    ordered_st = sorted(statuses.values(), key=lambda x: x.get("order", 99))
    headers = ["Jenis Apel", "Total Kegiatan"] + [s["name"] for s in ordered_st]
    rows = []
    for tid, v in agg.items():
        row = [types.get(tid, tid), len(v["events"])] + [v["by_key"].get(s["key"], 0) for s in ordered_st]
        rows.append(row)
    return headers, rows


@router.get("/attendance-recap/export")
async def export_recap(format: str = "excel", attendance_type_id: Optional[str] = None,
                       date_from: Optional[str] = None, date_to: Optional[str] = None,
                       user: dict = Depends(require("report.export", "attendance.export"))):
    headers, rows = await _recap_rows(attendance_type_id, date_from, date_to)
    title = "Rekap Kehadiran Apel"
    sub = f"Periode: {date_from or 'awal'} s/d {date_to or 'sekarang'}"
    if format == "pdf":
        return _stream(_pdf(title, sub, headers, rows, landscape_mode=True), "application/pdf", "rekap-kehadiran.pdf")
    return _stream(_excel(title, headers, rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "rekap-kehadiran.xlsx")


# ---------------- Combined APEL report (Informasi + Rekap + Notula + Dokumentasi) ----------------
@router.get("/apel/{event_id}/combined")
async def apel_combined(event_id: str, user: dict = Depends(require("report.view", "attendance.view"))):
    ev = await db.attendance_events.find_one({"id": event_id}, {"_id": 0})
    if not ev:
        raise HTTPException(404, "Kegiatan tidak ditemukan")
    from .attendance import compute_event_recap
    t = await db.attendance_types.find_one({"id": ev.get("attendance_type_id")}, {"_id": 0})
    ev["attendance_type_name"] = t["name"] if t else None
    recap = await compute_event_recap(event_id)
    note = await db.attendance_notes.find_one({"event_id": event_id}, {"_id": 0})
    docs = [{"title": d.get("title"), "file_name": d.get("file_name"), "note": d.get("note")} async for d in db.event_documents.find({"event_id": event_id}, {"_id": 0})]
    records = []
    statuses = {s["id"]: s async for s in db.attendance_statuses.find({})}
    async for r in db.attendance_records.find({"event_id": event_id}):
        emp = await db.employees.find_one({"id": r["employee_id"]}, {"_id": 0, "name": 1, "nip": 1})
        st = statuses.get(r.get("status_id"))
        records.append({"name": emp.get("name") if emp else "-", "nip": emp.get("nip") if emp else "-",
                        "status": st["name"] if st else "-", "keterangan": r.get("keterangan")})
    return {"event": ev, "recap": recap, "note": note, "documents": docs, "records": records}


@router.get("/apel/{event_id}/export")
async def apel_export(event_id: str, user: dict = Depends(require("report.export", "attendance.export"))):
    data = await apel_combined(event_id, user)
    ev = data["event"]
    recap = data["recap"]
    note = data["note"] or {}

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=16 * mm, leftMargin=16 * mm, rightMargin=16 * mm)
    styles = getSampleStyleSheet()
    h = ParagraphStyle("h", parent=styles["Title"], fontSize=14, textColor=NAVY, spaceAfter=2)
    org = ParagraphStyle("org", parent=styles["Normal"], fontSize=9, textColor=GOLD)
    sec = ParagraphStyle("sec", parent=styles["Heading2"], fontSize=11, textColor=NAVY, spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=9, leading=13)
    elems = [Paragraph("LAPAS GUNUNGSITOLI", org),
             Paragraph(f"LAPORAN KEGIATAN {ev.get('attendance_type_name','APEL').upper()} — {ev.get('session','')}", h),
             Paragraph(ev.get("code", ""), body)]

    elems.append(Paragraph("A. Informasi Kegiatan", sec))
    info = [["Tanggal", ev.get("date")], ["Waktu", ev.get("time")], ["Jenis Apel", ev.get("attendance_type_name")],
            ["Sesi", ev.get("session")], ["Lokasi", ev.get("location")], ["Pembina", ev.get("pembina")],
            ["Jabatan Pembina", ev.get("pembina_jabatan")], ["Catatan", ev.get("catatan")]]
    ti = Table([[Paragraph(f"<b>{a}</b>", body), Paragraph(str(b or "-"), body)] for a, b in info], colWidths=[45 * mm, 120 * mm])
    ti.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")), ("BACKGROUND", (0, 0), (0, -1), LIGHT)]))
    elems.append(ti)

    elems.append(Paragraph("B. Rekap Kehadiran", sec))
    rh = ["Total", "Hadir"] + [s["label"] for s in recap["by_status"] if s["key"] != "hadir"] + ["Persentase"]
    rv = [recap["total"], recap["present"]] + [s["count"] for s in recap["by_status"] if s["key"] != "hadir"] + [f"{recap['percentage']}%"]
    tr = Table([rh, rv])
    tr.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), NAVY), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    elems.append(tr)

    if data["records"]:
        elems.append(Paragraph("Daftar Kehadiran", sec))
        rows = [["Nama", "NIP", "Status", "Keterangan"]] + [[r["name"], r["nip"], r["status"], r["keterangan"] or "-"] for r in data["records"]]
        td = Table(rows, repeatRows=1, colWidths=[55 * mm, 35 * mm, 30 * mm, 45 * mm])
        td.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), NAVY), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")), ("FONTSIZE", (0, 0), (-1, -1), 8),
                                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT])]))
        elems.append(td)

    elems.append(Paragraph("C. Notula / Amanat", sec))
    for label, key in [("Tema/Judul", "tema"), ("Pokok-Pokok Amanat", "pokok"), ("Isi/Ringkasan", "isi"), ("Arahan/Tindak Lanjut", "tindak_lanjut"), ("Catatan Tambahan", "catatan")]:
        elems.append(Paragraph(f"<b>{label}:</b> {note.get(key) or '-'}", body))

    elems.append(Paragraph("D. Dokumentasi Kegiatan", sec))
    if data["documents"]:
        for d in data["documents"]:
            elems.append(Paragraph(f"• {d['title']} {('— ' + d['file_name']) if d.get('file_name') else ''} {('(' + d['note'] + ')') if d.get('note') else ''}", body))
    else:
        elems.append(Paragraph("Tidak ada dokumentasi.", body))

    elems.append(Spacer(1, 10))
    elems.append(Paragraph(f"Dicetak: {datetime.now().strftime('%d-%m-%Y %H:%M')}", ParagraphStyle("f", parent=body, fontSize=7, textColor=colors.grey)))
    doc.build(elems)
    buf.seek(0)
    return _stream(buf, "application/pdf", f"laporan-{ev.get('code','apel')}.pdf")
